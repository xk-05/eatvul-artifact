import csv
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from matplotlib import font_manager


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
FIG_DIR = REPORT_DIR / "figures"
OUT_DOCX = REPORT_DIR / "EaTVul_细粒度防御前后对比报告.docx"


PREVIOUS_CSV = ROOT / "results" / "eatvul_defense" / "lodo_calib_fpr_0.1_results.csv"
LOCAL_CSV = (
    ROOT
    / "results"
    / "eatvul_local_defense"
    / "localize_sanitize_w1536_s768_fpr0.01_max2_guided_benign_gate_no_sample_gate_results.csv"
)
CASCADE_CSV = (
    ROOT
    / "results"
    / "eatvul_local_defense"
    / "localize_sanitize_w1536_s768_fpr0.01_max2_guided_benign_gate_sample_gate0.5_results.csv"
)


DATASETS = ["asterisk", "openssl", "cwe119", "cwe399"]


def configure_matplotlib_fonts():
    candidates = [
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            font_manager.fontManager.addfont(str(path))
            name = font_manager.FontProperties(fname=str(path)).get_name()
            plt.rcParams["font.sans-serif"] = [name]
            plt.rcParams["font.family"] = "sans-serif"
            plt.rcParams["axes.unicode_minus"] = False
            return font_manager.FontProperties(fname=str(path))
    return None


def read_rows(path):
    with path.open(encoding="utf-8") as handle:
        return {row["dataset"]: row for row in csv.DictReader(handle)}


def f(row, key):
    return float(row[key])


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for margin, value in [("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")]:
                node = tc_mar.find(qn(f"w:{margin}"))
                if node is None:
                    node = OxmlElement(f"w:{margin}")
                    tc_mar.append(node)
                node.set(qn("w:w"), value)
                node.set(qn("w:type"), "dxa")


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.style = "Normal"
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        rest = text[len(bold_prefix) :]
        r = p.add_run(rest)
    else:
        r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(91, 103, 112)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def set_document_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, color in [
        ("Heading 1", 17, "17365D"),
        ("Heading 2", 13, "1F4D78"),
        ("Heading 3", 11, "365F91"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def plot_asr(previous, local, cascade):
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    x = np.arange(len(DATASETS))
    width = 0.18
    fig, ax = plt.subplots(figsize=(9.5, 4.8), dpi=180)
    series = [
        ("Baseline attack", [f(local[d], "baseline_asr") for d in DATASETS], "#5B6770", -1.5),
        ("Sample gate defense", [f(previous[d], "defended_asr") for d in DATASETS], "#2E74B5", -0.5),
        ("Local sanitize", [f(local[d], "sanitized_asr") for d in DATASETS], "#4C956C", 0.5),
        ("Cascade sanitize", [f(cascade[d], "sanitized_asr") for d in DATASETS], "#D08C30", 1.5),
    ]
    for label, values, color, offset in series:
        ax.bar(x + width * offset, values, width, label=label, color=color)
    ax.set_ylabel("Attack Success Rate (lower is better)")
    ax.set_xticks(x)
    ax.set_xticklabels(DATASETS)
    ax.set_ylim(0, 0.85)
    ax.grid(axis="y", alpha=0.25)
    ax.legend(ncol=4, frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.02), fontsize=8)
    ax.set_title("ASR before and after defense")
    fig.tight_layout()
    path = FIG_DIR / "asr_comparison.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_f1(previous, local, cascade):
    x = np.arange(len(DATASETS))
    width = 0.2
    fig, ax = plt.subplots(figsize=(9.5, 4.8), dpi=180)
    series = [
        ("Clean before", [f(local[d], "target_clean_f1") for d in DATASETS], "#5B6770", -1.5),
        ("Sample gate defense", [f(previous[d], "defended_clean_f1") for d in DATASETS], "#2E74B5", -0.5),
        ("Local sanitize", [f(local[d], "sanitized_clean_f1") for d in DATASETS], "#4C956C", 0.5),
        ("Cascade sanitize", [f(cascade[d], "sanitized_clean_f1") for d in DATASETS], "#D08C30", 1.5),
    ]
    for label, values, color, offset in series:
        ax.bar(x + width * offset, values, width, label=label, color=color)
    ax.set_ylabel("Clean F1 (higher is better)")
    ax.set_xticks(x)
    ax.set_xticklabels(DATASETS)
    ax.set_ylim(0, 1.0)
    ax.grid(axis="y", alpha=0.25)
    ax.legend(ncol=4, frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.02), fontsize=8)
    ax.set_title("Clean-set impact of each defense")
    fig.tight_layout()
    path = FIG_DIR / "clean_f1_comparison.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_flow(font_prop=None):
    fig, ax = plt.subplots(figsize=(9.5, 3.4), dpi=180)
    ax.axis("off")
    boxes = [
        ("1. 漏洞模型初判", "只对被判 benign\n的样本启动净化"),
        ("2. 局部窗口定位", "1536-token window\n768-token stride"),
        ("3. 异常特征打分", "稀有 token / 困惑度\n命名孤立 / AST 结构"),
        ("4. 模型反馈确认", "删除后 vulnerable\n概率上升才保留"),
        ("5. 删除并复检", "输出 sanitized AST-token\n再跑检测器"),
    ]
    xs = np.linspace(0.08, 0.92, len(boxes))
    for i, (title, body) in enumerate(boxes):
        x = xs[i]
        rect = plt.Rectangle((x - 0.075, 0.30), 0.15, 0.48, fill=True, color="#E8EEF5", ec="#1F4D78", lw=1.2)
        ax.add_patch(rect)
        ax.text(x, 0.62, title, ha="center", va="center", fontsize=8.5, weight="bold", fontproperties=font_prop)
        ax.text(x, 0.45, body, ha="center", va="center", fontsize=7.1, linespacing=1.25, fontproperties=font_prop)
        if i < len(boxes) - 1:
            ax.annotate("", xy=(xs[i + 1] - 0.085, 0.54), xytext=(x + 0.085, 0.54), arrowprops=dict(arrowstyle="->", lw=1.4))
    fig.tight_layout()
    path = FIG_DIR / "pipeline_flow.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def pct(value):
    return f"{value:.3f}"


def add_result_table(doc, previous, local, cascade):
    headers = [
        "数据集",
        "原始 ASR",
        "样本级 ASR",
        "局部净化 ASR",
        "级联净化 ASR",
        "原始 Clean F1",
        "样本级 Clean F1",
        "局部 Clean F1",
        "级联 Clean F1",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    style_table(table)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF", size=7.5)
        set_cell_shading(table.rows[0].cells[i], "17365D")
    for d in DATASETS:
        cells = table.add_row().cells
        values = [
            d,
            pct(f(local[d], "baseline_asr")),
            pct(f(previous[d], "defended_asr")),
            pct(f(local[d], "sanitized_asr")),
            pct(f(cascade[d], "sanitized_asr")),
            pct(f(local[d], "target_clean_f1")),
            pct(f(previous[d], "defended_clean_f1")),
            pct(f(local[d], "sanitized_clean_f1")),
            pct(f(cascade[d], "sanitized_clean_f1")),
        ]
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, size=7.5)
    return table


def add_metric_table(doc):
    rows = [
        ("ASR", "攻击成功率。ADV 样本被漏洞模型误判为 benign 的比例，越低表示防御越有效。"),
        ("Clean F1", "干净测试集上的漏洞检测 F1。用于衡量防御是否破坏正常检测能力。"),
        ("ADV modified rate", "ADV 样本中实际被定位并删除片段的比例，表示局部净化触发频率。"),
        ("ADV removed ratio", "被删除 token 数占原始 token 数的平均比例，用于观察净化强度。"),
        ("Prob gain", "删除候选窗口后，目标模型输出 vulnerable 概率的提升。它是模型反馈确认信号。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    style_table(table)
    set_cell_text(table.rows[0].cells[0], "指标", bold=True, color="FFFFFF")
    set_cell_text(table.rows[0].cells[1], "含义", bold=True, color="FFFFFF")
    set_cell_shading(table.rows[0].cells[0], "17365D")
    set_cell_shading(table.rows[0].cells[1], "17365D")
    for metric, meaning in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], metric)
        set_cell_text(cells[1], meaning)


def build():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    font_prop = configure_matplotlib_fonts()
    previous = read_rows(PREVIOUS_CSV)
    local = read_rows(LOCAL_CSV)
    cascade = read_rows(CASCADE_CSV)

    asr_fig = plot_asr(previous, local, cascade)
    f1_fig = plot_f1(previous, local, cascade)
    flow_fig = plot_flow(font_prop)

    doc = Document()
    set_document_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("EaTVul 细粒度插入片段定位与净化防御：前后对比报告")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("17365D")

    subtitle = doc.add_paragraph()
    r = subtitle.add_run("基于 AST-token 局部窗口、异常特征、模型反馈确认和删除复检的防御实现")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(91, 103, 112)

    add_heading(doc, "一、结论摘要", 1)
    add_body(
        doc,
        "本次改动把之前的“样本级前置拦截”推进为“局部定位后删除/净化”的防御流程。新脚本会在 AST-token 序列上定位疑似插入片段，删除后重新运行漏洞检测器，从而模拟论文级防御中更关键的定位与净化环节。",
    )
    add_body(
        doc,
        "主要效果：OpenSSL 数据集上，局部净化可将 ASR 从 0.760 降到 0.700，同时 clean F1 从 0.829 仅降到 0.819；加入样本级 gate 后 clean F1 保持 0.829，ASR 降到 0.720。整体看，前置样本级检测仍更强，但新方法已经具备“定位、删除、复检”的论文级防御骨架。",
    )

    add_heading(doc, "二、改动前后方法差异", 1)
    table = doc.add_table(rows=1, cols=4)
    style_table(table)
    for i, h in enumerate(["版本", "处理粒度", "防御动作", "优缺点"]):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[i], "17365D")
    rows = [
        ("原始 EaTVul", "无防御", "展示 benign 片段插入可以绕过漏洞检测器", "证明攻击有效，但没有识别或清除插入片段。"),
        ("样本级前置检测", "函数/样本级", "若样本像 ADV，则直接按 vulnerable 处理", "ASR 下降明显，但不能告诉我们插入片段在哪里。"),
        ("新增局部净化", "AST-token 窗口级", "定位疑似插入 span，删除后复检", "具备定位删除能力，但受限于没有源码/PDG真实边界。"),
        ("级联净化", "样本级 gate + 窗口级", "先判断是否疑似攻击，再做局部删除", "误伤更低，但召回和 ASR 降幅也更保守。"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)

    add_heading(doc, "三、防御流程图", 1)
    doc.add_picture(str(flow_fig), width=Inches(6.4))
    add_caption(doc, "图 1：细粒度定位与净化流程。异常特征只产生候选窗口，最终删除由目标模型概率反馈确认。")

    add_heading(doc, "四、步骤含义说明", 1)
    steps = [
        ("漏洞模型初判", "先运行目标漏洞检测器。因为 EaTVul 的攻击目标是把 vulnerable 样本伪装成 benign，所以默认只对初判为 benign 的样本启动净化。"),
        ("滑动窗口定位", "在 AST-token 序列上使用 1536 token 窗口和 768 token 步长扫描，得到候选局部片段。这个近似替代真实源码/PDG边界。"),
        ("异常特征打分", "每个窗口计算罕见 token、未见 token、局部 bigram NLL、孤立变量名、声明密度、函数定义密度、AST 结构密度等指标。"),
        ("模型反馈确认", "仅当删除候选窗口后，漏洞模型的 vulnerable 概率提升超过阈值，才认为该窗口可能是对抗插入片段。"),
        ("删除与复检", "删除被确认的 AST-token span，生成 sanitized JSONL，再将净化后的样本输入漏洞检测器，计算 ASR 和 clean F1。"),
    ]
    for title_text, body_text in steps:
        add_body(doc, f"{title_text}：{body_text}", bold_prefix=f"{title_text}：")

    add_heading(doc, "五、前后对比图", 1)
    doc.add_picture(str(asr_fig), width=Inches(6.4))
    add_caption(doc, "图 2：ASR 对比。ASR 越低，说明攻击越难绕过漏洞检测器。")
    doc.add_picture(str(f1_fig), width=Inches(6.4))
    add_caption(doc, "图 3：Clean F1 对比。该指标衡量防御是否损害正常漏洞检测能力。")

    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    landscape.top_margin = Inches(0.7)
    landscape.bottom_margin = Inches(0.7)
    landscape.left_margin = Inches(0.65)
    landscape.right_margin = Inches(0.65)
    add_heading(doc, "六、详细结果表", 1)
    add_result_table(doc, previous, local, cascade)

    add_heading(doc, "七、指标解释", 1)
    add_metric_table(doc)

    add_heading(doc, "八、效果分析", 1)
    add_body(
        doc,
        "样本级前置检测的 ASR 降幅最大，因为它不需要定位边界，只要判断整段代码像 ADV 就拦截。这适合安全网关式部署，但不满足“定位后删除/降权”的研究目标。",
    )
    add_body(
        doc,
        "局部净化实现了真正的片段级动作：它会输出 localized_spans，并生成删除后的 sanitized JSONL。OpenSSL 上效果最好，说明该数据集中插入片段具有较强跨项目/声明簇异常特征，窗口级定位能够恢复部分漏洞信号。",
    )
    add_body(
        doc,
        "Asterisk 上 clean F1 下降明显，说明仅用 AST-token 固定窗口会误删正常但异常分较高的声明片段。加入样本级 gate 可以降低误伤，但同时会漏掉部分 ADV，因此 ASR 降幅变小。",
    )
    add_body(
        doc,
        "CWE119 和 CWE399 上局部净化几乎没有降低 ASR，说明插入片段和原始函数在 AST-token 统计上更难区分，或者攻击成功依赖的不是单个高异常窗口。这也是下一步必须接源码边界或 PDG 的原因。",
    )

    add_heading(doc, "九、文件与复现实验命令", 1)
    add_body(doc, "新增脚本：scripts/eatvul_localize_sanitize.py")
    add_body(doc, "说明文档：LOCALIZATION_SANITIZATION_DEFENSE.md")
    add_body(doc, "主实验结果：results/eatvul_local_defense/localize_sanitize_w1536_s768_fpr0.01_max2_guided_benign_gate_no_sample_gate_results.csv")
    add_body(doc, "级联实验结果：results/eatvul_local_defense/localize_sanitize_w1536_s768_fpr0.01_max2_guided_benign_gate_sample_gate0.5_results.csv")

    add_heading(doc, "十、下一步建议", 1)
    add_body(
        doc,
        "当前实现已经具备论文级防御框架，但仍是 AST-token 近似。若要进一步提升为强论文结果，应恢复源码函数，使用 Joern 或 Tree-sitter 构建 AST/CFG/PDG，把固定窗口替换为完整语法子树或 PDG 孤立组件，再做删除或降权。这样可以避免 arbitrary token window 带来的误删，并让“控制流不可达概率、数据依赖孤立度”真正落地。",
    )

    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    print(build())
