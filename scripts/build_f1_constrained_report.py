import csv
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from matplotlib import font_manager


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
FIG_DIR = REPORT_DIR / "figures_f1_constrained"
OUT_DOCX = REPORT_DIR / "EaTVul_F1约束防御实验报告.docx"
RESULT_DIR = ROOT / "results" / "eatvul_f1_constrained_defense"
DATASETS = ["asterisk", "openssl", "cwe119", "cwe399"]
DROPS = [0.02, 0.03, 0.05]


def configure_fonts():
    for path in [Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/simhei.ttf")]:
        if path.exists():
            font_manager.fontManager.addfont(str(path))
            name = font_manager.FontProperties(fname=str(path)).get_name()
            plt.rcParams["font.sans-serif"] = [name]
            plt.rcParams["font.family"] = "sans-serif"
            plt.rcParams["axes.unicode_minus"] = False
            return font_manager.FontProperties(fname=str(path))
    return None


def result_path(drop):
    return RESULT_DIR / f"f1_constrained_maxf1drop{drop:g}_benignonly_results.csv"


def read_rows(path):
    with path.open(encoding="utf-8") as handle:
        return {row["dataset"]: row for row in csv.DictReader(handle)}


def f(rows, dataset, key):
    return float(rows[dataset][key])


def fmt(value):
    return f"{value:.3f}" if isinstance(value, float) else str(value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=8.3):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for margin, value in [("top", "80"), ("bottom", "80"), ("start", "100"), ("end", "100")]:
                node = tc_mar.find(qn(f"w:{margin}"))
                if node is None:
                    node = OxmlElement(f"w:{margin}")
                    tc_mar.append(node)
                node.set(qn("w:w"), value)
                node.set(qn("w:type"), "dxa")


def set_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color in [("Heading 1", 17, "17365D"), ("Heading 2", 13, "1F4D78")]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_body(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(91, 103, 112)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def plot_asr_reduction(all_rows):
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    x = np.arange(len(DATASETS))
    width = 0.24
    colors = ["#2E74B5", "#4C956C", "#D08C30"]
    fig, ax = plt.subplots(figsize=(9.5, 4.6), dpi=180)
    for idx, drop in enumerate(DROPS):
        values = [f(all_rows[drop], dataset, "asr_reduction") for dataset in DATASETS]
        ax.bar(x + (idx - 1) * width, values, width, label=f"F1 drop <= {drop:.2f}", color=colors[idx])
    ax.set_xticks(x)
    ax.set_xticklabels(DATASETS)
    ax.set_ylim(0, 0.40)
    ax.set_ylabel("ASR reduction (higher is better)")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(ncol=3, frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.02), fontsize=8)
    ax.set_title("ASR reduction under clean-F1 constraints")
    fig.tight_layout()
    path = FIG_DIR / "asr_reduction_by_budget.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_f1_drop(all_rows):
    x = np.arange(len(DATASETS))
    width = 0.24
    colors = ["#2E74B5", "#4C956C", "#D08C30"]
    fig, ax = plt.subplots(figsize=(9.5, 4.6), dpi=180)
    for idx, drop in enumerate(DROPS):
        values = [f(all_rows[drop], dataset, "clean_f1_drop") for dataset in DATASETS]
        ax.bar(x + (idx - 1) * width, values, width, label=f"Budget {drop:.2f}", color=colors[idx])
        ax.axhline(drop, color=colors[idx], linestyle="--", linewidth=0.8, alpha=0.45)
    ax.set_xticks(x)
    ax.set_xticklabels(DATASETS)
    ax.set_ylim(0, 0.06)
    ax.set_ylabel("Observed clean F1 drop")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(ncol=3, frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.02), fontsize=8)
    ax.set_title("Clean F1 degradation stays within the configured budget")
    fig.tight_layout()
    path = FIG_DIR / "f1_drop_by_budget.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_pipeline(font_prop=None):
    fig, ax = plt.subplots(figsize=(9.5, 3.1), dpi=180)
    ax.axis("off")
    boxes = [
        ("目标模型初判", "只处理被判 benign\n的可疑样本"),
        ("插入检测器打分", "LODO 训练\n输出攻击概率"),
        ("F1 约束校准", "搜索阈值\nClean F1 drop <= 预算"),
        ("防御动作", "可疑 benign\n改判为 vulnerable"),
        ("实验验证", "比较 Clean F1\n和 ASR 下降"),
    ]
    xs = np.linspace(0.08, 0.92, len(boxes))
    for i, (title, body) in enumerate(boxes):
        x = xs[i]
        rect = plt.Rectangle((x - 0.075, 0.30), 0.15, 0.48, fill=True, color="#E8EEF5", ec="#1F4D78", lw=1.2)
        ax.add_patch(rect)
        ax.text(x, 0.62, title, ha="center", va="center", fontsize=8.5, weight="bold", fontproperties=font_prop)
        ax.text(x, 0.45, body, ha="center", va="center", fontsize=7.2, linespacing=1.25, fontproperties=font_prop)
        if i < len(boxes) - 1:
            ax.annotate("", xy=(xs[i + 1] - 0.085, 0.54), xytext=(x + 0.085, 0.54), arrowprops=dict(arrowstyle="->", lw=1.4))
    fig.tight_layout()
    path = FIG_DIR / "f1_constrained_pipeline.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def result_table(doc, rows):
    headers = ["预算", "数据集", "Clean F1 前", "Clean F1 后", "F1下降", "ASR 前", "ASR 后", "ASR下降", "clean非漏洞FPR"]
    table = doc.add_table(rows=1, cols=len(headers))
    style_table(table)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF", size=7.4)
        set_cell_shading(table.rows[0].cells[i], "17365D")
    for drop in DROPS:
        for dataset in DATASETS:
            values = [
                f"{drop:.2f}",
                dataset,
                fmt(f(rows[drop], dataset, "target_clean_f1")),
                fmt(f(rows[drop], dataset, "defended_clean_f1")),
                fmt(f(rows[drop], dataset, "clean_f1_drop")),
                fmt(f(rows[drop], dataset, "baseline_asr")),
                fmt(f(rows[drop], dataset, "defended_asr")),
                fmt(f(rows[drop], dataset, "asr_reduction")),
                fmt(f(rows[drop], dataset, "clean_non_vul_fpr")),
            ]
            cells = table.add_row().cells
            for i, value in enumerate(values):
                set_cell_text(cells[i], value, size=7.4)


def build():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    font_prop = configure_fonts()
    rows = {drop: read_rows(result_path(drop)) for drop in DROPS}
    asr_fig = plot_asr_reduction(rows)
    f1_fig = plot_f1_drop(rows)
    pipeline_fig = plot_pipeline(font_prop)

    doc = Document()
    set_styles(doc)
    title = doc.add_paragraph()
    r = title.add_run("EaTVul F1 约束防御实验报告")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("17365D")
    sub = doc.add_paragraph()
    r = sub.add_run("目标：在保证 clean F1 的前提下降低攻击成功率 ASR")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(91, 103, 112)

    add_heading(doc, "一、防御方法")
    add_body(doc, "本次实现的是 F1-constrained adversarial gate。它复用前面的插入片段检测器，但不再固定使用 0.5 或人工阈值，而是在 clean 数据上自动搜索阈值，使防御后的 clean F1 下降不超过给定预算。")
    add_body(doc, "防御动作只作用于目标漏洞模型初判为 benign 的样本：如果该样本同时被插入检测器判为高风险，则把最终结果改判为 vulnerable。这样符合 EaTVul 的攻击目标，也避免对已判 vulnerable 的样本做无意义扰动。")

    add_heading(doc, "二、流程图")
    doc.add_picture(str(pipeline_fig), width=Inches(6.4))
    add_caption(doc, "图 1：F1 约束防御流程。核心是阈值不手工设定，而是在 clean F1 预算内自动校准。")

    add_heading(doc, "三、实验设置")
    add_body(doc, "目标模型使用 TF-IDF + Logistic Regression 作为可快速复现实验的漏洞检测器；插入检测器使用 leave-one-dataset-out 训练，即测试某个数据集时，检测器只用其他三个数据集的 clean/ADV 训练。")
    add_body(doc, "实验测试三档 clean F1 最大下降预算：0.02、0.03、0.05。每档预算都输出 clean F1 前后变化、ASR 前后变化、ADV recall 和 clean 非漏洞误报率。")

    add_heading(doc, "四、对比图")
    doc.add_picture(str(asr_fig), width=Inches(6.4))
    add_caption(doc, "图 2：在不同 clean F1 预算下的 ASR 降幅。OpenSSL 收益最明显，CWE119 在 0.05 预算下也有可观下降。")
    doc.add_picture(str(f1_fig), width=Inches(6.4))
    add_caption(doc, "图 3：实际 clean F1 下降均没有超过对应预算，说明约束阈值选择生效。")

    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    landscape.top_margin = Inches(0.7)
    landscape.bottom_margin = Inches(0.7)
    landscape.left_margin = Inches(0.55)
    landscape.right_margin = Inches(0.55)
    add_heading(doc, "五、详细结果表")
    result_table(doc, rows)

    add_heading(doc, "六、结论")
    add_body(doc, "可行性结论：在 clean F1 下降不超过 0.02 的严格条件下，OpenSSL 的 ASR 从 0.760 降到 0.560，下降 0.200；clean F1 从 0.829 到 0.810，下降 0.020，满足约束。")
    add_body(doc, "当预算放宽到 0.05 时，OpenSSL ASR 降到 0.400，CWE119 ASR 从 0.720 降到 0.670。Asterisk 和 CWE399 上收益有限，说明检测器对这些数据集的 ADV 分数与 clean 分数可分性不足。")
    add_body(doc, "因此，该防御方法在部分数据集上能在保证 F1 的情况下显著降低 ASR；但跨数据集普适性还需要更强的特征或源码/PDG 级边界来提升。")

    add_heading(doc, "七、复现命令")
    add_body(doc, "conda run -n eatvul python scripts\\eatvul_f1_constrained_defense.py --datasets asterisk openssl cwe119 cwe399 --max-clean-f1-drop 0.02")
    add_body(doc, "结果目录：results\\eatvul_f1_constrained_defense")
    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    print(build())
