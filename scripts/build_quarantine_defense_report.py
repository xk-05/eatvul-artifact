import csv
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from matplotlib import font_manager


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
FIG_DIR = REPORT_DIR / "figures_quarantine"
OUT_DOCX = REPORT_DIR / "EaTVul_强ASR降低阻断防御报告.docx"
RESULT_DIR = ROOT / "results" / "eatvul_quarantine_defense"
DATASETS = ["asterisk", "openssl", "cwe119", "cwe399"]
BUDGETS = [0.10, 0.20, 0.30]


def configure_fonts():
    for path in [Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/simhei.ttf")]:
        if path.exists():
            font_manager.fontManager.addfont(str(path))
            name = font_manager.FontProperties(fname=str(path)).get_name()
            plt.rcParams["font.sans-serif"] = [name]
            plt.rcParams["font.family"] = "sans-serif"
            plt.rcParams["axes.unicode_minus"] = False
            return font_manager.FontProperties(fname=str(path))
    return None


def result_path(budget):
    return RESULT_DIR / f"quarantine_cleanblock{budget:g}_benignonly_results.csv"


def read_rows(path):
    with path.open(encoding="utf-8") as handle:
        return {row["dataset"]: row for row in csv.DictReader(handle)}


def f(rows, dataset, key):
    return float(rows[dataset][key])


def fmt(value):
    return f"{value:.3f}" if isinstance(value, float) else str(value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=7.4):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for margin, value in [("top", "80"), ("bottom", "80"), ("start", "90"), ("end", "90")]:
                node = tc_mar.find(qn(f"w:{margin}"))
                if node is None:
                    node = OxmlElement(f"w:{margin}")
                    tc_mar.append(node)
                node.set(qn("w:w"), value)
                node.set(qn("w:type"), "dxa")


def set_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color in [("Heading 1", 17, "17365D"), ("Heading 2", 13, "1F4D78")]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_body(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(91, 103, 112)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def plot_asr(rows_by_budget):
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    x = np.arange(len(DATASETS))
    width = 0.18
    fig, ax = plt.subplots(figsize=(9.5, 4.7), dpi=180)
    baseline = [f(rows_by_budget[0.10], d, "baseline_asr") for d in DATASETS]
    ax.bar(x - 1.5 * width, baseline, width, label="Baseline attack", color="#5B6770")
    colors = ["#2E74B5", "#4C956C", "#D08C30"]
    for i, budget in enumerate(BUDGETS):
        values = [f(rows_by_budget[budget], d, "quarantine_asr") for d in DATASETS]
        ax.bar(x + (i - 0.5) * width, values, width, label=f"Block budget {budget:.0%}", color=colors[i])
    ax.set_xticks(x)
    ax.set_xticklabels(DATASETS)
    ax.set_ylim(0, 0.85)
    ax.set_ylabel("Residual ASR after quarantine (lower is better)")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(ncol=4, frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.02), fontsize=8)
    ax.set_title("Quarantine defense sharply lowers residual ASR")
    fig.tight_layout()
    path = FIG_DIR / "quarantine_asr.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_cost(rows_by_budget):
    x = np.arange(len(DATASETS))
    width = 0.24
    colors = ["#2E74B5", "#4C956C", "#D08C30"]
    fig, ax = plt.subplots(figsize=(9.5, 4.7), dpi=180)
    for i, budget in enumerate(BUDGETS):
        values = [f(rows_by_budget[budget], d, "clean_block_rate") for d in DATASETS]
        ax.bar(x + (i - 1) * width, values, width, label=f"Block budget {budget:.0%}", color=colors[i])
    ax.set_xticks(x)
    ax.set_xticklabels(DATASETS)
    ax.set_ylim(0, 0.32)
    ax.set_ylabel("Observed clean block rate")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(ncol=3, frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.02), fontsize=8)
    ax.set_title("Usability cost: clean samples sent to review")
    fig.tight_layout()
    path = FIG_DIR / "quarantine_clean_cost.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_pipeline(font_prop=None):
    fig, ax = plt.subplots(figsize=(9.5, 3.1), dpi=180)
    ax.axis("off")
    boxes = [
        ("目标模型初判", "只看被判 benign\n的放行风险"),
        ("插入检测器", "输出 ADV 风险分数\nLODO 训练"),
        ("阻断阈值校准", "按 clean 阻断预算\n选择阈值"),
        ("Quarantine", "高风险 benign\n不直接放行"),
        ("ASR 评估", "被阻断 ADV\n不算攻击成功"),
    ]
    xs = np.linspace(0.08, 0.92, len(boxes))
    for i, (title, body) in enumerate(boxes):
        x = xs[i]
        rect = plt.Rectangle((x - 0.075, 0.30), 0.15, 0.48, fill=True, color="#E8EEF5", ec="#1F4D78", lw=1.2)
        ax.add_patch(rect)
        ax.text(x, 0.62, title, ha="center", va="center", fontsize=8.5, weight="bold", fontproperties=font_prop)
        ax.text(x, 0.45, body, ha="center", va="center", fontsize=7.2, linespacing=1.25, fontproperties=font_prop)
        if i < len(boxes) - 1:
            ax.annotate("", xy=(xs[i + 1] - 0.085, 0.54), xytext=(x + 0.085, 0.54), arrowprops=dict(arrowstyle="->", lw=1.4))
    fig.tight_layout()
    path = FIG_DIR / "quarantine_pipeline.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def result_table(doc, rows_by_budget):
    headers = ["预算", "数据集", "Clean F1", "ASR 前", "ASR 后", "ASR下降", "ADV阻断率", "Clean阻断率", "Clean非漏洞阻断率"]
    table = doc.add_table(rows=1, cols=len(headers))
    style_table(table)
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[i], "17365D")
    for budget in BUDGETS:
        for dataset in DATASETS:
            row = rows_by_budget[budget][dataset]
            values = [
                f"{budget:.0%}",
                dataset,
                fmt(float(row["target_clean_f1"])),
                fmt(float(row["baseline_asr"])),
                fmt(float(row["quarantine_asr"])),
                fmt(float(row["asr_reduction"])),
                fmt(float(row["adv_block_rate"])),
                fmt(float(row["clean_block_rate"])),
                fmt(float(row["clean_non_vul_block_rate"])),
            ]
            cells = table.add_row().cells
            for i, value in enumerate(values):
                set_cell_text(cells[i], value)


def build():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    font_prop = configure_fonts()
    rows = {budget: read_rows(result_path(budget)) for budget in BUDGETS}
    asr_fig = plot_asr(rows)
    cost_fig = plot_cost(rows)
    pipeline_fig = plot_pipeline(font_prop)

    doc = Document()
    set_styles(doc)
    title = doc.add_paragraph()
    r = title.add_run("EaTVul 强 ASR 降低：Quarantine 阻断防御报告")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("17365D")
    sub = doc.add_paragraph()
    r = sub.add_run("目标：保持漏洞模型 clean F1，同时显著降低攻击绕过成功率")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(91, 103, 112)

    add_heading(doc, "一、防御方法")
    add_body(doc, "在前一个 F1 约束改判方法上，本次把防御动作改成 quarantine 阻断：目标模型判为 benign 且插入检测器风险高的样本不直接放行，而是进入复核/二阶段重检队列。")
    add_body(doc, "这与生产环境的安全网关更一致：clean F1 仍由原漏洞检测器给出，不因为强制改判产生大量 false positive；ASR 计算时，被阻断的 ADV 不再算成功绕过。代价则单独报告为 clean block rate。")

    add_heading(doc, "二、流程图")
    doc.add_picture(str(pipeline_fig), width=Inches(6.4))
    add_caption(doc, "图 1：阻断式防御流程。核心是用 clean 阻断预算控制可用性代价，同时最大化阻断 ADV。")

    add_heading(doc, "三、前后对比图")
    doc.add_picture(str(asr_fig), width=Inches(6.4))
    add_caption(doc, "图 2：残余 ASR。10% 预算已显著降低 Asterisk/OpenSSL，20%-30% 是高安全模式。")
    doc.add_picture(str(cost_fig), width=Inches(6.4))
    add_caption(doc, "图 3：clean 阻断代价。F1 保持原值，但需要把一部分 clean benign 样本送入复核。")

    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    landscape.top_margin = Inches(0.7)
    landscape.bottom_margin = Inches(0.7)
    landscape.left_margin = Inches(0.55)
    landscape.right_margin = Inches(0.55)
    add_heading(doc, "四、详细结果表")
    result_table(doc, rows)

    add_heading(doc, "五、结论")
    add_body(doc, "10% clean 非漏洞阻断预算下，Asterisk ASR 从 0.640 降到 0.360，OpenSSL ASR 从 0.760 降到 0.120；clean F1 保持原目标模型结果。")
    add_body(doc, "20% 预算下，OpenSSL ASR 进一步降到 0.040，Asterisk 降到 0.320，CWE399 也从 0.695 降到 0.535。")
    add_body(doc, "30% 高安全预算下，Asterisk ASR 降到 0.240，OpenSSL 维持 0.040，CWE119 降到 0.540。该模式显著降低 ASR，但需要接受更高 clean 阻断率。")
    add_body(doc, "因此，如果目标是显著降低 ASR，并允许可疑 benign 进入复核而不是直接给二分类标签，quarantine 防御比强制改判更合适。")

    add_heading(doc, "六、复现命令")
    add_body(doc, "conda run -n eatvul python scripts\\eatvul_quarantine_defense.py --datasets asterisk openssl cwe119 cwe399 --clean-block-budget 0.10")
    add_body(doc, "结果目录：results\\eatvul_quarantine_defense")
    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    print(build())
