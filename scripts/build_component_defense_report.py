import csv
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from matplotlib import font_manager


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
FIG_DIR = REPORT_DIR / "figures_component"
OUT_DOCX = REPORT_DIR / "EaTVul_结构组件级防御前后对比报告.docx"

SAMPLE_CSV = ROOT / "results" / "eatvul_defense" / "lodo_calib_fpr_0.1_results.csv"
WINDOW_CSV = (
    ROOT
    / "results"
    / "eatvul_local_defense"
    / "localize_sanitize_w1536_s768_fpr0.01_max2_guided_benign_gate_no_sample_gate_results.csv"
)
COMPONENT_CSV = (
    ROOT
    / "results"
    / "eatvul_component_defense"
    / "component_sanitize_fpr0.01_cluster3_max1_benign_gate_no_sample_gate_results.csv"
)

DATASETS = ["asterisk", "openssl", "cwe119", "cwe399"]


def configure_fonts():
    for path in [Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/simhei.ttf")]:
        if path.exists():
            font_manager.fontManager.addfont(str(path))
            name = font_manager.FontProperties(fname=str(path)).get_name()
            plt.rcParams["font.sans-serif"] = [name]
            plt.rcParams["font.family"] = "sans-serif"
            plt.rcParams["axes.unicode_minus"] = False
            return font_manager.FontProperties(fname=str(path))
    return None


def read_rows(path):
    with path.open(encoding="utf-8") as handle:
        return {row["dataset"]: row for row in csv.DictReader(handle)}


def f(rows, dataset, key):
    return float(rows[dataset][key])


def fmt(value):
    return f"{value:.3f}" if isinstance(value, float) else str(value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row in table.rows:
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = tc_pr.first_child_found_in("w:tcMar")
            if tc_mar is None:
                tc_mar = OxmlElement("w:tcMar")
                tc_pr.append(tc_mar)
            for margin, value in [("top", "80"), ("bottom", "80"), ("start", "100"), ("end", "100")]:
                node = tc_mar.find(qn(f"w:{margin}"))
                if node is None:
                    node = OxmlElement(f"w:{margin}")
                    tc_mar.append(node)
                node.set(qn("w:w"), value)
                node.set(qn("w:type"), "dxa")


def set_styles(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.85)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for name, size, color in [("Heading 1", 17, "17365D"), ("Heading 2", 13, "1F4D78")]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_body(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(91, 103, 112)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def plot_asr(sample, window, component):
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    x = np.arange(len(DATASETS))
    width = 0.2
    series = [
        ("Baseline attack", [f(window, d, "baseline_asr") for d in DATASETS], "#5B6770", -1.5),
        ("Sample-level gate", [f(sample, d, "defended_asr") for d in DATASETS], "#2E74B5", -0.5),
        ("Fixed-window sanitize", [f(window, d, "sanitized_asr") for d in DATASETS], "#4C956C", 0.5),
        ("AST-component sanitize", [f(component, d, "component_asr") for d in DATASETS], "#D08C30", 1.5),
    ]
    fig, ax = plt.subplots(figsize=(9.5, 4.8), dpi=180)
    for label, values, color, offset in series:
        ax.bar(x + offset * width, values, width, label=label, color=color)
    ax.set_xticks(x)
    ax.set_xticklabels(DATASETS)
    ax.set_ylim(0, 0.85)
    ax.set_ylabel("Attack Success Rate (lower is better)")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(ncol=4, frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.02), fontsize=8)
    ax.set_title("ASR before and after each defense")
    fig.tight_layout()
    path = FIG_DIR / "component_asr_comparison.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_f1(sample, window, component):
    x = np.arange(len(DATASETS))
    width = 0.2
    series = [
        ("Clean before", [f(window, d, "target_clean_f1") for d in DATASETS], "#5B6770", -1.5),
        ("Sample-level gate", [f(sample, d, "defended_clean_f1") for d in DATASETS], "#2E74B5", -0.5),
        ("Fixed-window sanitize", [f(window, d, "sanitized_clean_f1") for d in DATASETS], "#4C956C", 0.5),
        ("AST-component sanitize", [f(component, d, "component_clean_f1") for d in DATASETS], "#D08C30", 1.5),
    ]
    fig, ax = plt.subplots(figsize=(9.5, 4.8), dpi=180)
    for label, values, color, offset in series:
        ax.bar(x + offset * width, values, width, label=label, color=color)
    ax.set_xticks(x)
    ax.set_xticklabels(DATASETS)
    ax.set_ylim(0, 1.0)
    ax.set_ylabel("Clean F1 (higher is better)")
    ax.grid(axis="y", alpha=0.25)
    ax.legend(ncol=4, frameon=False, loc="upper center", bbox_to_anchor=(0.5, 1.02), fontsize=8)
    ax.set_title("Clean-set impact of each defense")
    fig.tight_layout()
    path = FIG_DIR / "component_f1_comparison.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_pipeline(font_prop=None):
    fig, ax = plt.subplots(figsize=(9.5, 3.2), dpi=180)
    ax.axis("off")
    boxes = [
        ("源码/结构输入", "优先 Tree-sitter/Joern\n当前数据用 AST-token"),
        ("组件边界恢复", "top-level AST 子树\n相邻组件簇"),
        ("PDG 孤立近似", "组件内外标识符\n重叠越少越孤立"),
        ("异常+反馈排序", "结构异常分\n+ vulnerable 概率增益"),
        ("组件级删除复检", "删除完整组件\n再跑漏洞检测"),
    ]
    xs = np.linspace(0.08, 0.92, len(boxes))
    for i, (title, body) in enumerate(boxes):
        x = xs[i]
        rect = plt.Rectangle((x - 0.075, 0.30), 0.15, 0.48, fill=True, color="#E8EEF5", ec="#1F4D78", lw=1.2)
        ax.add_patch(rect)
        ax.text(x, 0.62, title, ha="center", va="center", fontsize=8.4, weight="bold", fontproperties=font_prop)
        ax.text(x, 0.45, body, ha="center", va="center", fontsize=7.1, linespacing=1.25, fontproperties=font_prop)
        if i < len(boxes) - 1:
            ax.annotate("", xy=(xs[i + 1] - 0.085, 0.54), xytext=(x + 0.085, 0.54), arrowprops=dict(arrowstyle="->", lw=1.4))
    fig.tight_layout()
    path = FIG_DIR / "component_pipeline.png"
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def result_table(doc, sample, window, component):
    headers = [
        "数据集",
        "Baseline ASR",
        "样本级 ASR",
        "窗口 ASR",
        "组件 ASR",
        "Clean F1 before",
        "样本级 F1",
        "窗口 F1",
        "组件 F1",
        "组件 ADV 修改率",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    style_table(table)
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF", size=7.4)
        set_cell_shading(table.rows[0].cells[i], "17365D")
    for d in DATASETS:
        values = [
            d,
            fmt(f(window, d, "baseline_asr")),
            fmt(f(sample, d, "defended_asr")),
            fmt(f(window, d, "sanitized_asr")),
            fmt(f(component, d, "component_asr")),
            fmt(f(window, d, "target_clean_f1")),
            fmt(f(sample, d, "defended_clean_f1")),
            fmt(f(window, d, "sanitized_clean_f1")),
            fmt(f(component, d, "component_clean_f1")),
            fmt(f(component, d, "adv_modified_rate")),
        ]
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, size=7.4)


def build():
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    font_prop = configure_fonts()
    sample = read_rows(SAMPLE_CSV)
    window = read_rows(WINDOW_CSV)
    component = read_rows(COMPONENT_CSV)
    asr_fig = plot_asr(sample, window, component)
    f1_fig = plot_f1(sample, window, component)
    pipeline_fig = plot_pipeline(font_prop)

    doc = Document()
    set_styles(doc)
    title = doc.add_paragraph()
    r = title.add_run("EaTVul 源码/结构组件级防御：前后对比报告")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string("17365D")
    sub = doc.add_paragraph()
    r = sub.add_run("从固定窗口净化升级到完整 AST 组件/PDG 孤立近似的定位删除")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(91, 103, 112)

    add_heading(doc, "一、实现结论")
    add_body(doc, "本次新增 scripts/eatvul_component_sanitize.py，将定位单位从固定 token 窗口升级为结构组件：根据 AST-token 中的结构节点和 depth 数字恢复 top-level 语法子树，并构造相邻组件簇。删除时不再切任意窗口，而是删除完整语法组件。")
    add_body(doc, "由于当前公开资源没有原始 C/C++ 源码，也没有 Joern/Tree-sitter 依赖，本实现采用 AST-token 组件边界作为源码 AST/PDG 的可复现实验近似；脚本和报告中明确保留了后续接入 Tree-sitter/Joern 的替换点。")

    add_heading(doc, "二、新流程图")
    doc.add_picture(str(pipeline_fig), width=Inches(6.4))
    add_caption(doc, "图 1：结构组件级防御流程。当前数据走 AST-token 组件边界；若补齐源码，可替换为 Tree-sitter/Joern 子树和 PDG 组件。")

    add_heading(doc, "三、关键步骤含义")
    for text in [
        "组件边界恢复：使用 FUNCTION_DEF、SIMPLE_DECL、CLASS_DEF、VAR_DECL 等 top-level 结构节点和 depth=1 标记切出完整语法组件，避免固定窗口把语义结构切碎。",
        "相邻组件簇：EaTVul 插入往往是一串声明或辅助函数，因此候选不仅包含单个组件，也包含连续组件簇。本实验正式参数使用 max-cluster=3。",
        "PDG 孤立近似：在没有真实 PDG 的情况下，用组件内标识符与组件外标识符的重叠率近似数据依赖。重叠越少，越像独立插入的无害代码。",
        "模型反馈排序：候选组件必须同时满足异常分阈值和删除后 vulnerable 概率上升，最终按 probability gain 选择删除对象。",
        "复检：删除完整组件后重新运行漏洞检测器，比较 ASR 和 clean F1，评估防御收益与误伤。",
    ]:
        add_body(doc, text)

    add_heading(doc, "四、前后对比图")
    doc.add_picture(str(asr_fig), width=Inches(6.4))
    add_caption(doc, "图 2：ASR 越低越好。样本级拦截仍最强，组件级净化具备更真实的定位删除语义。")
    doc.add_picture(str(f1_fig), width=Inches(6.4))
    add_caption(doc, "图 3：Clean F1 越高越好。组件级方法在 OpenSSL/CWE119/CWE399 上保持较稳，但 Asterisk 仍存在误删。")

    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    landscape.top_margin = Inches(0.7)
    landscape.bottom_margin = Inches(0.7)
    landscape.left_margin = Inches(0.55)
    landscape.right_margin = Inches(0.55)
    add_heading(doc, "五、详细对比表")
    result_table(doc, sample, window, component)

    add_heading(doc, "六、效果分析")
    add_body(doc, "样本级前置检测在 ASR 上仍最强，因为它不需要定位边界，只要判定样本像 ADV 就直接拦截。但它不能删除或降权具体片段。")
    add_body(doc, "固定窗口净化可以定位并删除局部 token span，但边界不是语法/PDG 单元，容易切到正常代码。")
    add_body(doc, "组件级净化更接近论文级防御：删除的是完整 AST 组件或组件簇，并引入数据依赖孤立度近似。OpenSSL/CWE119/CWE399 上 clean F1 基本保持；Asterisk 上 clean F1 下降，说明 AST-token 组件仍不能完全替代真实源码/PDG。")
    add_body(doc, "组件级 ASR 降幅有限，主要原因是公开数据缺少原始源码和真实 PDG，导致控制流不可达概率、数据依赖孤立度只能近似。真正的强论文版本应接入 Joern 或 Tree-sitter，从源码构建 AST/CFG/PDG 后删除完整语法子树或 PDG 孤立组件。")

    add_heading(doc, "七、复现命令")
    add_body(doc, "组件级正式实验命令：conda run -n eatvul python scripts\\eatvul_component_sanitize.py --datasets asterisk openssl cwe119 cwe399 --calibrate-component-fpr 0.01 --max-cluster 3 --max-spans 1 --candidate-components 16 --min-prob-gain 0.001 --gate-on-benign --calibration-items 200")
    add_body(doc, "结果文件：results\\eatvul_component_defense\\component_sanitize_fpr0.01_cluster3_max1_benign_gate_no_sample_gate_results.csv")

    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    print(build())
